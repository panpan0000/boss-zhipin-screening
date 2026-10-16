#!/usr/bin/env python3
"""把 zhipin_extract.py 抓到的原始 JSON 清洗成最终交付物。

清洗动作：
  1. 剔除混在列表里的 VIP 推广卡片（无名无优势）
  2. 按 (姓名, 年龄, 薪资) 去重 —— 无限滚动会重复渲染
  3. 输出 clean JSON / Markdown / Excel（含统计 sheet）

用法:
    python3 build_outputs.py zhipin_candidates.json --prefix zhipin_candidates
"""
import json
import re
import time
from collections import Counter

# openpyxl 只在导出 Excel 时才需要，放函数内惰性导入，
# 这样 clean() / build_md() 在没有 openpyxl 的解释器里也能用。
SALARY_RE = re.compile(r"(\d+)\s*-\s*(\d+)\s*K")


def salary_mid(s):
    """取薪资区间中值，用于排序/筛选；解析不出返回 None。"""
    if not s:
        return None
    m = SALARY_RE.search(s)
    if m:
        return (int(m.group(1)) + int(m.group(2))) / 2
    return None


def build_md(data, cards):
    L = [f"# BOSS直聘 推荐牛人\n",
         f"- 目标职位: **{data.get('job_title') or '(页面未显示)'}**",
         f"- 候选人数: **{len(cards)}**",
         f"- 抓取时间: {time.strftime('%Y-%m-%d %H:%M:%S')}\n",
         "| # | 姓名 | 年龄 | 经验 | 学历 | 期望薪资 | 城市 | 期望职位 | 求职状态 | 活跃度 |",
         "|---|---|---|---|---|---|---|---|---|---|"]
    for c in cards:
        L.append("| {idx} | {name} | {age} | {exp} | {deg} | {sal} | {city} | {pos} | {st} | {act} |".format(
            idx=c["idx"], name=c["name"], age=c.get("age") or "-",
            exp=c.get("exp_text") or "-", deg=c.get("degree") or "-",
            sal=c.get("salary") or "-", city=c.get("expect_city") or "-",
            pos=c.get("expect_position") or "-",
            st=c.get("job_status") or "-", act=c.get("active") or "-"))

    L.append("\n## 逐人详情\n")
    for c in cards:
        L.append(f"### {c['idx']}. {c['name']} — {c.get('salary') or '薪资面议'}")
        L.append(f"- **基本**: {c.get('age')}岁 / {c.get('exp_text') or '-'} / "
                 f"{c.get('degree') or '-'} / {c.get('job_status') or '-'} / {c.get('active') or '-'}")
        L.append(f"- **期望**: {c.get('expect_city') or '-'} {c.get('expect_position') or ''}")
        if c.get("tags"):
            L.append(f"- **标签**: {' · '.join(c['tags'])}")
        if c.get("work_exps"):
            L.append(f"- **经历**: {' ；'.join(c['work_exps'])}")
        if c.get("edu_exps"):
            L.append(f"- **教育**: {' ；'.join(c['edu_exps'])}")
        if c.get("advantage"):
            L.append(f"- **优势**: {c['advantage']}")
        L.append("")
    return "\n".join(L)


def build_xlsx(cards, path):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "候选人"
    HEAD = ["#", "姓名", "年龄", "经验", "学历", "期望薪资", "薪资中值(K)", "城市", "期望职位",
            "求职状态", "活跃度", "应届", "标签", "工作经历", "教育经历", "优势描述"]
    ws.append(HEAD)
    fill = PatternFill("solid", fgColor="1F4E79")
    for i in range(1, len(HEAD) + 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    for c in cards:
        ws.append([
            c["idx"], c["name"], c.get("age"), c.get("exp_text"), c.get("degree"),
            c.get("salary"), salary_mid(c.get("salary")), c.get("expect_city"),
            c.get("expect_position"), c.get("job_status"), c.get("active"),
            "是" if c.get("is_fresh_grad") else "",
            " · ".join(c.get("tags") or []),
            " ；".join(c.get("work_exps") or []),
            " ；".join(c.get("edu_exps") or []),
            c.get("advantage") or "",
        ])

    for i, w in enumerate([5, 11, 6, 8, 7, 11, 12, 8, 15, 15, 11, 6, 26, 32, 30, 70], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wrap_cols = {13, 14, 15, 16}
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in wrap_cols)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEAD))}{ws.max_row}"

    # ---- 统计 sheet ----
    ws2 = wb.create_sheet("统计")
    ws2.append(["维度", "取值", "人数"])
    for h in range(1, 4):
        ws2.cell(row=1, column=h).font = Font(bold=True)
        ws2.cell(row=1, column=h).fill = fill
        ws2.cell(row=1, column=h).font = Font(bold=True, color="FFFFFF")

    for label, key in [("学历", "degree"), ("城市", "expect_city"),
                       ("期望职位", "expect_position"), ("求职状态", "job_status")]:
        for val, n in Counter(c.get(key) for c in cards if c.get(key)).most_common(15):
            ws2.append([label, val, n])

    ages = [c["age"] for c in cards if c.get("age")]
    exps = [c["exp_years"] for c in cards if c.get("exp_years") is not None]
    fresh = sum(1 for c in cards if c.get("is_fresh_grad"))
    if ages:
        ws2.append(["年龄", f"平均 {sum(ages)/len(ages):.1f} 岁（{min(ages)}–{max(ages)}）", len(ages)])
    if exps:
        ws2.append(["经验", f"平均 {sum(exps)/len(exps):.1f} 年（{min(exps)}–{max(exps)}）", len(exps)])
    ws2.append(["应届生", f"{fresh} 人（{fresh*100//max(len(cards),1)}%）", fresh])

    for col, w in zip("ABC", (12, 46, 8)):
        ws2.column_dimensions[col].width = w
    wb.save(path)
